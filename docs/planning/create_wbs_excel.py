import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "WBS"

headers = ["Level", "WBS", "Task Description", "Assigned To", "Start", "End", "Notes", "Chi phí"]
ws.append(headers)

# Styling for header
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center
    cell.border = thin_border

# Data array fully populated with Notes and Costs
data = [
    [1, "1", "EVManager - HTQL Trung tâm Tiệc cưới", "Hoàng (PM)", "07/09/2026", "01/11/2026", "Tổng toàn dự án", "93,500,000"],
    [2, "1.1", "Khởi tạo dự án", "Hoàng", "07/09/2026", "13/09/2026", "GĐ1: Định hình dự án", "5,000,000"],
    [3, "1.1.1", "Lập Tôn chỉ dự án (Project Charter)", "Hoàng", "07/09/2026", "09/09/2026", "Lấy chữ ký các bên", "2,000,000"],
    [3, "1.1.2", "Lập Phát biểu phạm vi (Scope Statement)", "Hiển", "10/09/2026", "11/09/2026", "Chốt ranh giới scope", "2,000,000"],
    [3, "1.1.3", "Họp Kick-off và phân công đội ngũ", "All", "12/09/2026", "13/09/2026", "Biên bản họp Kick-off", "1,000,000"],
    
    [2, "1.2", "Lập kế hoạch dự án", "Hoàng", "14/09/2026", "20/09/2026", "GĐ2: Lập KH chi tiết", "5,000,000"],
    [3, "1.2.1", "Xây dựng WBS", "Hoàng", "14/09/2026", "15/09/2026", "File Excel/WPP", "1,500,000"],
    [3, "1.2.2", "Lập tiến độ (Gantt Chart)", "Hoàng", "16/09/2026", "17/09/2026", "Sử dụng MS Project", "1,500,000"],
    [3, "1.2.3", "Lập ngân sách (Budget Plan)", "Hiển", "18/09/2026", "19/09/2026", "Duyệt chi phí 93.5M", "1,000,000"],
    [3, "1.2.4", "Lập kế hoạch quản lý rủi ro", "Hoàng", "20/09/2026", "20/09/2026", "Risk Register Log", "1,000,000"],
    
    [2, "1.3", "Phân tích và Thiết kế", "Hiển, Nhân, Hậu", "21/09/2026", "04/10/2026", "GĐ3: Thiết kế", "20,000,000"],
    [3, "1.3.1", "Đặc tả yêu cầu nghiệp vụ (SRS)", "Hiển", "21/09/2026", "25/09/2026", "10-14 Use cases", "7,000,000"],
    [3, "1.3.2", "Thiết kế CSDL (ERD, Data Dictionary)", "Hậu", "26/09/2026", "29/09/2026", "8-12 bảng", "4,000,000"],
    [3, "1.3.3", "Thiết kế giao diện (UI/UX Prototype)", "Nhân", "27/09/2026", "02/10/2026", "Figma link", "5,000,000"],
    [3, "1.3.4", "Thiết kế API & Kiến trúc hệ thống", "Phúc", "30/09/2026", "04/10/2026", "Swagger/Postman", "4,000,000"],
    
    [2, "1.4", "Phát triển phần mềm (Coding)", "Phúc, Nhân", "05/10/2026", "18/10/2026", "GĐ4: Lập trình", "40,000,000"],
    [3, "1.4.1", "Thiết lập môi trường (Dev/Staging)", "Phúc", "05/10/2026", "06/10/2026", "Setup Server/CI-CD", "3,000,000"],
    [3, "1.4.2", "Lập trình Backend & Database", "Phúc, Hậu", "07/10/2026", "18/10/2026", "Logic & API", "15,000,000"],
    [3, "1.4.3", "Lập trình Frontend (Web & Desktop)", "Nhân", "07/10/2026", "18/10/2026", "UI integration", "15,000,000"],
    [3, "1.4.4", "Tích hợp hệ thống (Integration)", "Hoàng, Phúc", "15/10/2026", "18/10/2026", "Ghép API - Frontend", "7,000,000"],
    
    [2, "1.5", "Kiểm thử (Testing)", "Hậu, All", "19/10/2026", "25/10/2026", "GĐ5: Quality Control", "13,500,000"],
    [3, "1.5.1", "Xây dựng kịch bản kiểm thử (Test Cases)", "Hậu", "19/10/2026", "20/10/2026", "Tối thiểu 35 TCs", "3,500,000"],
    [3, "1.5.2", "Kiểm thử chức năng và hiệu năng", "Hậu", "21/10/2026", "23/10/2026", "Log bug lên Jira", "4,000,000"],
    [3, "1.5.3", "Kiểm thử chấp nhận người dùng (UAT)", "Hiển, Hoàng", "24/10/2026", "24/10/2026", "Nghiệm thu với khách", "2,000,000"],
    [3, "1.5.4", "Ghi nhận và sửa lỗi (Fix bugs)", "Nhân, Phúc", "22/10/2026", "25/10/2026", "Fix severity 1,2", "4,000,000"],
    
    [2, "1.6", "Triển khai và Đóng dự án", "Hoàng", "26/10/2026", "01/11/2026", "GĐ6: Đóng dự án", "10,000,000"],
    [3, "1.6.1", "Triển khai lên Production (Golive)", "Phúc", "26/10/2026", "27/10/2026", "Live server", "4,000,000"],
    [3, "1.6.2", "Soạn tài liệu Hướng dẫn sử dụng", "Hiển", "28/10/2026", "29/10/2026", "PDF Manual", "2,000,000"],
    [3, "1.6.3", "Nghiệm thu và Bàn giao hệ thống", "Hoàng, All", "30/10/2026", "31/10/2026", "Ký biên bản nghiệm thu", "2,000,000"],
    [3, "1.6.4", "Họp rút kinh nghiệm (Lessons Learned)", "All", "01/11/2026", "01/11/2026", "Đóng repo, lưu tài liệu", "2,000,000"]
]

phase_fill = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")

for row_idx, row_data in enumerate(data, 2):
    for col_idx, cell_value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
        cell.border = thin_border
        
        # Color phases differently (Level 1 and 2)
        if row_data[0] == 1 or row_data[0] == 2:
            cell.fill = phase_fill
            cell.font = Font(bold=True)
            
        if col_idx in [1, 2, 5, 6]:
            cell.alignment = align_center
            
        # Align Chi phí right
        if col_idx == 8:
            cell.alignment = align_right

# Adjust column widths
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 45
ws.column_dimensions['D'].width = 18
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 25 # Increased for Notes
ws.column_dimensions['H'].width = 15

wb.save("d:\\EVManager\\docs\\planning\\WBS_EVManager_Format_v2.xlsx")
