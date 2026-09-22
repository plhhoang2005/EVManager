import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "WBS_Detailed"

headers = ["Level", "WBS", "Task Description", "Assigned To", "Start", "End", "Notes", "Chi phí (VNĐ)"]
ws.append(headers)

# Styling for header
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center
    cell.border = thin_border

# Data arrays: [Level, WBS, Task, Assigned, Start, End, Note, Cost]
data = [
    [1, "1", "EVManager - HTQL Trung tâm Hội nghị Tiệc cưới", "All", "07/09/2026", "01/11/2026", "Tổng toàn dự án", "93,500,000"],
    
    [2, "1.1", "Khởi tạo dự án", "Hoàng", "07/09/2026", "11/09/2026", "Giai đoạn 1", "5,000,000"],
    [3, "1.1.1", "Lập Tôn chỉ dự án (Project Charter)", "Hoàng", "07/09/2026", "08/09/2026", "Phê duyệt Charter", "1,500,000"],
    [3, "1.1.2", "Lập Phát biểu phạm vi (Scope Statement)", "Hiển", "09/09/2026", "10/09/2026", "Chốt scope", "2,000,000"],
    [3, "1.1.3", "Khảo sát nguồn lực & Họp Kick-off", "Hoàng, All", "11/09/2026", "11/09/2026", "Họp team", "1,500,000"],
    
    [2, "1.2", "Lập kế hoạch dự án", "Hoàng", "12/09/2026", "18/09/2026", "Giai đoạn 2", "5,000,000"],
    [3, "1.2.1", "Xây dựng Sơ đồ WBS", "Hoàng", "12/09/2026", "13/09/2026", "Cấu trúc công việc", "1,000,000"],
    [3, "1.2.2", "Lập tiến độ (Gantt Chart)", "Hoàng", "14/09/2026", "15/09/2026", "Cột mốc thời gian", "1,000,000"],
    [3, "1.2.3", "Lập ngân sách (Cost Baseline)", "Hiển", "16/09/2026", "16/09/2026", "Kế hoạch chi phí", "1,000,000"],
    [3, "1.2.4", "Lập kế hoạch Quản lý Rủi ro", "Hoàng", "17/09/2026", "17/09/2026", "Risk Matrix", "1,000,000"],
    [3, "1.2.5", "Lập kế hoạch Kiểm thử (Test Plan)", "Hậu", "18/09/2026", "18/09/2026", "Quy trình test", "1,000,000"],
    
    [2, "1.3", "Phân tích và Thiết kế", "Hiển, Nhân, Hậu, Phúc", "19/09/2026", "04/10/2026", "Giai đoạn 3", "20,000,000"],
    [3, "1.3.1", "Đặc tả yêu cầu nghiệp vụ (SRS)", "Hiển", "19/09/2026", "23/09/2026", "Phân tích hệ thống", "6,000,000"],
    [4, "1.3.1.1", "Xác định Actor & Danh sách Use Cases", "Hiển", "19/09/2026", "20/09/2026", "Actor: Admin, Sales...", "1,500,000"],
    [4, "1.3.1.2", "Đặc tả Use Case: Đặt sảnh & Hợp đồng", "Hiển", "21/09/2026", "22/09/2026", "Luồng cốt lõi", "2,500,000"],
    [4, "1.3.1.3", "Đặc tả Use Case: Thanh toán & Báo cáo", "Hiển", "23/09/2026", "23/09/2026", "Kế toán, Thống kê", "2,000,000"],
    [3, "1.3.2", "Thiết kế Cơ sở Dữ liệu", "Hậu", "24/09/2026", "26/09/2026", "Kiến trúc DB", "4,000,000"],
    [4, "1.3.2.1", "Vẽ biểu đồ quan hệ thực thể (ERD)", "Hậu", "24/09/2026", "25/09/2026", "Mô hình quan hệ", "2,000,000"],
    [4, "1.3.2.2", "Lập từ điển dữ liệu (Data Dictionary)", "Hậu", "26/09/2026", "26/09/2026", "Schema 10-12 bảng", "2,000,000"],
    [3, "1.3.3", "Thiết kế Giao diện (UI/UX)", "Nhân", "27/09/2026", "01/10/2026", "Thiết kế hình ảnh", "6,000,000"],
    [4, "1.3.3.1", "Vẽ Wireframe (Bố cục chức năng)", "Nhân", "27/09/2026", "28/09/2026", "Khung xương UI", "2,000,000"],
    [4, "1.3.3.2", "Thiết kế Prototype (Figma)", "Nhân", "29/09/2026", "30/09/2026", "Thiết kế chi tiết", "3,000,000"],
    [4, "1.3.3.3", "Review và Chốt giao diện", "Nhân, Hiển", "01/10/2026", "01/10/2026", "Lấy xác nhận KH", "1,000,000"],
    [3, "1.3.4", "Thiết kế Kiến trúc & API", "Phúc", "02/10/2026", "04/10/2026", "Base code", "4,000,000"],
    [4, "1.3.4.1", "Thiết kế System Architecture", "Phúc", "02/10/2026", "03/10/2026", "Kiến trúc tổng", "2,000,000"],
    [4, "1.3.4.2", "Viết tài liệu API (Swagger/Postman)", "Phúc, Hiển", "04/10/2026", "04/10/2026", "API spec", "2,000,000"],
    
    [2, "1.4", "Phát triển phần mềm (Coding)", "Phúc, Nhân", "05/10/2026", "20/10/2026", "Giai đoạn 4", "40,000,000"],
    [3, "1.4.1", "Thiết lập môi trường", "Phúc", "05/10/2026", "06/10/2026", "Hạ tầng", "3,000,000"],
    [4, "1.4.1.1", "Mua Domain & Thuê Cloud Server", "Phúc, Hoàng", "05/10/2026", "05/10/2026", "Vercel / VPS", "1,500,000"],
    [4, "1.4.1.2", "Cấu hình Database & CI/CD Pipeline", "Phúc, Hậu", "06/10/2026", "06/10/2026", "Deploy tự động", "1,500,000"],
    [3, "1.4.2", "Lập trình Backend (API)", "Phúc", "07/10/2026", "15/10/2026", "Logic xử lý", "15,000,000"],
    [4, "1.4.2.1", "Module Auth (Đăng nhập, RBAC)", "Phúc", "07/10/2026", "08/10/2026", "Bảo mật", "2,000,000"],
    [4, "1.4.2.2", "Module Quản lý Sảnh & Dịch vụ", "Phúc", "09/10/2026", "10/10/2026", "CRUD cơ bản", "3,000,000"],
    [4, "1.4.2.3", "Module Đặt tiệc & Xử lý trùng lịch", "Phúc", "11/10/2026", "13/10/2026", "Thuật toán Check lịch", "5,000,000"],
    [4, "1.4.2.4", "Module Hợp đồng & Thanh toán", "Phúc", "14/10/2026", "15/10/2026", "Xử lý dòng tiền", "3,000,000"],
    [4, "1.4.2.5", "Module Thống kê & Báo cáo", "Phúc", "15/10/2026", "15/10/2026", "Dashboard", "2,000,000"],
    [3, "1.4.3", "Lập trình Frontend", "Nhân", "07/10/2026", "17/10/2026", "Giao diện Web", "15,000,000"],
    [4, "1.4.3.1", "Cắt HTML/CSS/JS theo UI/UX", "Nhân", "07/10/2026", "10/10/2026", "Dựng khung", "5,000,000"],
    [4, "1.4.3.2", "Ghép API: Auth & Dashboard", "Nhân", "11/10/2026", "12/10/2026", "Đăng nhập", "2,000,000"],
    [4, "1.4.3.3", "Ghép API: Lịch sảnh & Quản lý tiệc", "Nhân", "13/10/2026", "15/10/2026", "Trang chính", "5,000,000"],
    [4, "1.4.3.4", "Ghép API: Hợp đồng & Thanh toán", "Nhân", "16/10/2026", "17/10/2026", "Nghiệp vụ phụ", "3,000,000"],
    [3, "1.4.4", "Tích hợp Hệ thống (Integration)", "Hoàng, Phúc, Nhân", "18/10/2026", "20/10/2026", "Kiểm tra kết nối", "7,000,000"],
    
    [2, "1.5", "Kiểm thử (QC & Testing)", "Hậu", "21/10/2026", "27/10/2026", "Giai đoạn 5", "13,500,000"],
    [3, "1.5.1", "Xây dựng kịch bản kiểm thử (TCs)", "Hậu", "21/10/2026", "22/10/2026", "Viết TCs", "3,500,000"],
    [4, "1.5.1.1", "Viết TC cho Luồng chức năng", "Hậu", "21/10/2026", "21/10/2026", "Functional TC", "2,000,000"],
    [4, "1.5.1.2", "Viết TC Giao diện & Hiệu năng", "Hậu", "22/10/2026", "22/10/2026", "Non-functional", "1,500,000"],
    [3, "1.5.2", "Thực thi kiểm thử nội bộ", "Hậu", "23/10/2026", "24/10/2026", "Chạy test", "4,000,000"],
    [4, "1.5.2.1", "Test các module độc lập", "Hậu", "23/10/2026", "23/10/2026", "Unit Test", "2,000,000"],
    [4, "1.5.2.2", "Test luồng tích hợp hệ thống", "Hậu", "24/10/2026", "24/10/2026", "System Test", "2,000,000"],
    [3, "1.5.3", "Ghi nhận và Sửa lỗi (Fix bugs)", "Phúc, Nhân", "25/10/2026", "26/10/2026", "Sửa lỗi", "4,000,000"],
    [4, "1.5.3.1", "Log bugs lên Jira", "Hậu", "25/10/2026", "25/10/2026", "List bugs", "1,000,000"],
    [4, "1.5.3.2", "Dev sửa lỗi (Fix bugs)", "Phúc, Nhân", "25/10/2026", "26/10/2026", "Sửa severity cao", "2,000,000"],
    [4, "1.5.3.3", "Kiểm thử lại (Retest)", "Hậu", "26/10/2026", "26/10/2026", "Xác nhận fix", "1,000,000"],
    [3, "1.5.4", "Kiểm thử chấp nhận người dùng (UAT)", "Hiển, Hoàng", "27/10/2026", "27/10/2026", "Demo Khách", "2,000,000"],
    
    [2, "1.6", "Triển khai và Đóng dự án", "Hoàng", "28/10/2026", "01/11/2026", "Giai đoạn 6", "10,000,000"],
    [3, "1.6.1", "Triển khai lên Production (Golive)", "Phúc", "28/10/2026", "28/10/2026", "Release", "3,000,000"],
    [3, "1.6.2", "Bàn giao tài liệu", "Hiển", "29/10/2026", "30/10/2026", "Tài liệu", "3,000,000"],
    [4, "1.6.2.1", "Soạn HDSD (User Manual)", "Hiển", "29/10/2026", "29/10/2026", "Bản PDF", "1,500,000"],
    [4, "1.6.2.2", "Đóng gói Source code & Kỹ thuật", "Phúc", "30/10/2026", "30/10/2026", "Bàn giao kỹ thuật", "1,500,000"],
    [3, "1.6.3", "Nghiệm thu và Đóng dự án", "Hoàng", "31/10/2026", "01/11/2026", "Hoàn thành", "4,000,000"],
    [4, "1.6.3.1", "Ký biên bản nghiệm thu (Sign-off)", "Hoàng", "31/10/2026", "31/10/2026", "Chốt dự án", "2,000,000"],
    [4, "1.6.3.2", "Họp rút kinh nghiệm (Lessons Learned)","All", "01/11/2026", "01/11/2026", "Họp tổng kết", "2,000,000"]
]

phase_fill_12 = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid") # Darker blue for phases
phase_fill_3 = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Lighter blue for level 3

for row_idx, row_data in enumerate(data, 2):
    for col_idx, cell_value in enumerate(row_data, 1):
        # Indent Task Description based on Level for visual hierarchy
        if col_idx == 3:
            lvl = row_data[0]
            prefix = "  " * (lvl - 1)
            cell = ws.cell(row=row_idx, column=col_idx, value=prefix + str(cell_value))
        else:
            cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
            
        cell.border = thin_border
        
        # Color phases differently
        if row_data[0] in [1, 2]:
            cell.fill = phase_fill_12
            cell.font = Font(bold=True)
        elif row_data[0] == 3:
            cell.fill = phase_fill_3
            cell.font = Font(bold=True, italic=True)
            
        if col_idx in [1, 2, 5, 6]:
            cell.alignment = align_center
            
        # Align Chi phí right
        if col_idx == 8:
            cell.alignment = align_right

# Adjust column widths
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 12
ws.column_dimensions['C'].width = 50
ws.column_dimensions['D'].width = 18
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 25
ws.column_dimensions['H'].width = 18

wb.save("d:\\EVManager\\docs\\planning\\WBS_EVManager_Detailed.xlsx")
