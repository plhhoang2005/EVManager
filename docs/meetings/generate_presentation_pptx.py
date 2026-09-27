import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    NAVY = RGBColor(27, 54, 93)      # #1B365D
    GOLD = RGBColor(180, 130, 20)    # #B48214
    DARK = RGBColor(40, 50, 60)      # #28323C
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(245, 247, 250)
    CARD_BG = RGBColor(255, 255, 255)
    BORDER_COLOR = RGBColor(220, 225, 230)
    BLUE_ACCENT = RGBColor(14, 116, 180)
    RED_ACCENT = RGBColor(190, 40, 40)
    GREEN_ACCENT = RGBColor(30, 130, 60)

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="BÁO CÁO TIẾN ĐỘ GIỮA KỲ (MỐC 50%)"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = NAVY
        p1.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    t_box1 = slide1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(3.8))
    tf1 = t_box1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "DỰ ÁN HỆ THỐNG QUẢN LÝ TRUNG TÂM HỘI NGHỊ TIỆC CƯỚI"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(212, 175, 55) # Gold
    p.space_after = Pt(10)

    p = tf1.add_paragraph()
    p.text = "EVManager System - Midterm Progress & 50% Demo"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(14)

    p = tf1.add_paragraph()
    p.text = "Báo cáo tiến độ Tuần 1-3 | Phân tích WBS & EVM Day 20 | Demo 50% Tính Năng Cốt Lõi"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(200, 215, 230)
    p.space_after = Pt(28)

    p = tf1.add_paragraph()
    p.text = "Nhóm Thực Hiện: Phạm Lê Huy Hoàng (PM), Huỳnh Tấn Thọ, Võ Hoàng Nhân, Bùi Nguyễn Chí Hậu, Dương Khắc Đạt\nThời gian báo cáo: Ngày 20 (Cuối tuần 3 - 27/09/2026)"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(180, 200, 220)

    # -------------------------------------------------------------
    # SLIDE 2: Objectives & Problem Statement
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Mục Tiêu Dự Án & Bài Toán Nghiệp Vụ Thực Tế")
    
    # Left Card: Problem
    card1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
    card1.fill.solid()
    card1.fill.fore_color.rgb = RGBColor(253, 245, 245)
    card1.line.color.rgb = RGBColor(240, 200, 200)
    tf = card1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "THÁCH THỨC NGHIỆP VỤ THỰC TẾ"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(12)
    
    problems = [
        "Trùng lịch đặt sảnh: Do quản lý bằng sổ tay hoặc nhiều file Excel rời rạc giữa các bộ phận.",
        "Thất thoát phụ thu dịch vụ: Khó kiểm soát phát sinh giờ phát tiệc, rượu ngoại, âm thanh ánh sáng ngoài gói.",
        "Báo giá thủ công chậm trễ: Nhân viên mất 20-30 phút tính toán giá theo từng ca (sáng/tối), ngày thường/cuối tuần.",
        "Thiếu báo cáo dòng tiền: Quản lý khó nắm bắt các đợt thanh toán cọc và công nợ thực tế của từng hợp đồng."
    ]
    for prob in problems:
        p = tf.add_paragraph()
        p.text = "• " + prob
        p.font.size = Pt(12)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # Right Card: Solution
    card2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(245, 250, 245)
    card2.line.color.rgb = RGBColor(200, 230, 200)
    tf = card2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "GIẢI PHÁP HỆ THỐNG EVMANAGER"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(12)

    solutions = [
        "Khóa trùng lịch tự động: Kiểm tra tính sẵn sàng của sảnh theo ngày, ca sáng/tối với độ chính xác 100%.",
        "Bộ công cụ tính giá tự động (Pricing Calculator): Dự toán trọn gói sảnh, thực đơn và dịch vụ trong 10 giây.",
        "Chuẩn hóa quy trình 3NF (13 bảng): Quản lý chi tiết hợp đồng, phân rã thanh toán cọc 3 đợt rõ ràng.",
        "Dashboard chỉ số tài chính thời gian thực: Cập nhật tỷ lệ lấp đầy sảnh và doanh thu dự kiến tức thì."
    ]
    for sol in solutions:
        p = tf.add_paragraph()
        p.text = "✔ " + sol
        p.font.size = Pt(12)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: WBS & Budget Structure
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Cơ Cấu Phân Rã Công Việc & Ngân Sách Cơ Sở (BAC = 130 Triệu)")

    # Table of WBS
    rows, cols = 8, 5
    left, top, width, height = Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2)
    table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.columns[0].width = Inches(1.2)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(2.2)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(2.0)

    headers = ["Mã WBS", "Tên Gói Công Việc", "Ngân Sách (VNĐ)", "Tỷ Trọng", "Trạng Thái Day 20"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    wbs_data = [
        ["1.1", "Khởi động dự án & Thiết lập chuẩn hóa", "7,000,000 đ", "5.38%", "Hoàn thành 100%"],
        ["1.2", "Phân tích yêu cầu nghiệp vụ tiệc cưới", "13,000,000 đ", "10.00%", "Hoàn thành 100%"],
        ["1.3", "Thiết kế hệ thống & CSDL 13 bảng 3NF", "16,500,000 đ", "12.69%", "Hoàn thành 100%"],
        ["1.4", "Phát triển Backend Spring Boot 3", "30,000,000 đ", "23.08%", "Hoàn thành 87% SP1"],
        ["1.5", "Phát triển Frontend React + Vite SPA", "28,500,000 đ", "21.92%", "Hoàn thành 83% SP1"],
        ["1.6", "Kiểm thử hệ thống & Đảm bảo chất lượng", "18,000,000 đ", "13.85%", "Sprint 2 (Tuần 4-7)"],
        ["1.7", "Triển khai Staging & Đóng dự án", "17,000,000 đ", "13.08%", "Sprint 2 (Tuần 7-8)"]
    ]

    for row_idx, row in enumerate(wbs_data, start=1):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if row_idx % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = DARK
                if col_idx in [0, 2, 3, 4]:
                    p.alignment = PP_ALIGN.CENTER
                if col_idx == 4:
                    p.font.bold = True
                    if "Hoàn thành" in text:
                        p.font.color.rgb = GREEN_ACCENT

    # -------------------------------------------------------------
    # SLIDE 4: CPM & Fast-tracking
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Tiến Độ CPM, Nhận Diện Nút Thắt & Kế Hoạch Fast-tracking")

    # Box 1: CPM Path
    cpm_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.5))
    cpm_box.fill.solid()
    cpm_box.fill.fore_color.rgb = LIGHT_BG
    cpm_box.line.color.rgb = BORDER_COLOR
    tf = cpm_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ĐƯỜNG GĂNG DỰ ÁN (CRITICAL PATH - TỔNG THỜI GIAN 56 NGÀY / 8 TUẦN)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    
    p = tf.add_paragraph()
    p.text = "1.1 Khởi động (W1) ➔ 1.2 Phân tích (W1-W2) ➔ 1.3 Thiết kế 3NF (W2) ➔ 1.4-1.5 Phát triển (W3-W5) ➔ 1.6 Kiểm thử (W6-W7) ➔ 1.7 Triển khai (W8)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_before = Pt(8)

    # Box 2: Issue Identified
    c1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.2), Inches(5.7), Inches(3.7))
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(255, 250, 245)
    c1.line.color.rgb = RGBColor(240, 210, 180)
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PHÁT HIỆN NÚT THẮT TẠI TUẦN 2"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(190, 80, 20)
    p.space_after = Pt(8)

    items1 = [
        "9 công việc có nguy cơ trễ 24 person-hours vào cuối tuần 2.",
        "Nguyên nhân kỹ thuật: Đội ngũ quyết định nâng cấp CSDL từ 8 bảng lên 13 bảng chuẩn 3NF để bao quát nghiệp vụ thanh toán nhiều đợt và phân ca sáng/tối.",
        "Đường cong tiếp cận Spring Boot 3 và Flyway migration mất thêm thời gian nghiên cứu của nhóm Backend."
    ]
    for it in items1:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(6)

    # Box 3: Fast-tracking Execution
    c2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.2), Inches(5.7), Inches(3.7))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(245, 252, 245)
    c2.line.color.rgb = RGBColor(190, 230, 190)
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "GIẢI PHÁP FAST-TRACKING TUẦN 3 ĐÃ THỰC HIỆN"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(8)

    items2 = [
        "Song song hóa Frontend & Backend: Chốt API Contract sớm; Frontend phát triển trên Vite Mock Server không bị block.",
        "Huy động 48 giờ làm thêm (Overtime) của 5 thành viên trong Tuần 3 để hoàn thiện Schema DDL và bộ Mock API.",
        "Kết quả đo lường: Kéo chỉ số SPI từ 0.88 lên 0.944, sẵn sàng hoàn toàn cho buổi demo giữa kỳ 50% hôm nay."
    ]
    for it in items2:
        p = tf.add_paragraph()
        p.text = "✔ " + it
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: EVM Analysis Day 20
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Đo Lường Quản Trị Giá Trị Thu Được (EVM) Tại Day 20")

    # 4 Metric Cards
    metrics = [
        ("KẾ HOẠCH (PV)", "58,500,000 đ", "45.0% BAC", NAVY),
        ("THU ĐƯỢC (EV)", "55,250,000 đ", "42.5% BAC", BLUE_ACCENT),
        ("THỰC CHI (AC)", "63,300,000 đ", "48.7% BAC", RGBColor(180, 80, 20)),
        ("HIỆU SUẤT TIẾN ĐỘ", "SPI = 0.944", "SV = -3.25M VNĐ", GREEN_ACCENT)
    ]
    for idx, (m_title, m_val, m_sub, m_col) in enumerate(metrics):
        x = Inches(0.8 + idx * 2.98)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.5), Inches(2.8), Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        tf = card.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(120, 130, 140)
        p.alignment = PP_ALIGN.CENTER
        
        p = tf.add_paragraph()
        p.text = m_val
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = m_col
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(8)

        p = tf.add_paragraph()
        p.text = m_sub
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(4)

    # Detailed Explanation Table
    rows, cols = 5, 4
    table_shape = slide5.shapes.add_table(rows, cols, Inches(0.8), Inches(3.6), Inches(11.7), Inches(3.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(2.0)
    table.columns[3].width = Inches(5.5)

    evm_rows = [
        ["Chỉ Số Đo Lường", "Giá Trị", "Ngưỡng Đánh Giá", "Phân Tích Chi Tiết"],
        ["Phương sai tiến độ (SV)", "-3,250,000 đ", "SV >= -5.0M", "Chậm nhẹ ~1.5 ngày làm việc do mở rộng CSDL lên 13 bảng."],
        ["Phương sai chi phí (CV)", "-8,050,000 đ", "CV >= -10.0M", "Vượt chi phí do OT và chi phí nghiên cứu công nghệ mới ban đầu."],
        ["Chỉ số tiến độ (SPI)", "0.944", "SPI >= 0.90", "Tiến độ đạt 94.4% kế hoạch cơ sở (vùng kiểm soát tốt)."],
        ["Chỉ số chi phí (CPI)", "0.873", "CPI >= 0.85", "Cứ 1 đồng chi tiêu mang lại 0.873 đồng giá trị khối lượng."]
    ]
    for r_idx, row in enumerate(evm_rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = NAVY
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(11)
                    p.font.bold = True
                    p.font.color.rgb = WHITE
                    p.alignment = PP_ALIGN.CENTER
            else:
                cell.fill.fore_color.rgb = RGBColor(248, 250, 252) if r_idx % 2 == 0 else WHITE
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(11)
                    p.font.color.rgb = DARK
                    if c_idx in [0, 1, 2]:
                        p.alignment = PP_ALIGN.CENTER
                        p.font.bold = True

    # -------------------------------------------------------------
    # SLIDE 6: Forecasting & TCPI
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Dự Báo Ngân Sách Khi Hoàn Thành (EAC) & Chỉ Số TCPI")

    # Left: Forecasting Scenarios
    scen_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.3))
    scen_box.fill.solid()
    scen_box.fill.fore_color.rgb = WHITE
    scen_box.line.color.rgb = BORDER_COLOR
    tf = scen_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "CÁC KỊCH BẢN DỰ BÁO NGÂN SÁCH HOÀN THÀNH (EAC)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)

    scens = [
        ("Kịch bản 1: Giữ nguyên xu hướng hiện tại (Typical)", "EAC = BAC / CPI = 148,911,000 đ (VAC = -18.91M đ)\nXảy ra nếu nhóm tiếp tục gặp khó khăn kỹ thuật ở Sprint 2."),
        ("Kịch bản 2: Phần còn lại chạy theo định mức (Atypical)", "EAC = AC + (BAC - EV) = 138,050,000 đ (VAC = -8.05M đ)\nXảy ra khi các công việc Sprint 2 thực hiện đúng dự toán ban đầu."),
        ("Kịch bản 3: Tác động kép tiến độ & chi phí (Composite)", "EAC = AC + [(BAC - EV) / (CPI x SPI)] = 154,058,000 đ (VAC = -24.06M đ)\nKịch bản xấu nhất nếu trễ cả tiến độ và đội thêm chi phí."),
        ("Kịch bản 4: Có can thiệp quản trị phục hồi (Target)", "EAC = 132,500,000 đ (VAC = -2.50M đ - NẰM TRONG BUFFER)\nNhờ tái sử dụng 13 bảng CSDL & UI có sẵn, CPI Sprint 2 phục hồi >= 1.08.")
    ]
    for s_title, s_desc in scens:
        p = tf.add_paragraph()
        p.text = "▶ " + s_title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT
        p.space_before = Pt(6)
        
        p = tf.add_paragraph()
        p.text = s_desc
        p.font.size = Pt(10)
        p.font.color.rgb = DARK
        p.space_after = Pt(4)

    # Right: TCPI
    tcpi_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.5), Inches(4.6), Inches(5.3))
    tcpi_box.fill.solid()
    tcpi_box.fill.fore_color.rgb = RGBColor(245, 248, 255)
    tcpi_box.line.color.rgb = RGBColor(200, 220, 245)
    tf = tcpi_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "CHỈ SỐ HIỆU SUẤT CẦN ĐẠT (TCPI)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(12)

    p = tf.add_paragraph()
    p.text = "TCPI (Theo BAC gốc 130M):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DARK
    
    p = tf.add_paragraph()
    p.text = "1.121"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RED_ACCENT
    p.space_after = Pt(4)

    p = tf.add_paragraph()
    p.text = "Phần công việc còn lại phải đạt hiệu suất chi phí 112.1% để không vượt tổng ngân sách 130 triệu đồng."
    p.font.size = Pt(10)
    p.font.color.rgb = DARK
    p.space_after = Pt(14)

    p = tf.add_paragraph()
    p.text = "TCPI (Có tính quỹ dự phòng 10% = 143M):"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = DARK

    p = tf.add_paragraph()
    p.text = "0.938"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(4)

    p = tf.add_paragraph()
    p.text = "Khi sử dụng Quỹ dự phòng khẩn cấp đã duyệt trong COCOMO, áp lực giảm xuống còn 0.938 (rất khả thi)."
    p.font.size = Pt(10)
    p.font.color.rgb = DARK

    # -------------------------------------------------------------
    # SLIDE 7: Technical Architecture & 13-Table 3NF Database
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Kiến Trúc Kỹ Thuật Hiện Đại & CSDL 13 Bảng Chuẩn 3NF")

    # 3 Column Cards
    col_w = Inches(3.64)
    gap = Inches(0.39)

    # Col 1: Frontend
    fe_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), col_w, Inches(5.3))
    fe_card.fill.solid()
    fe_card.fill.fore_color.rgb = WHITE
    fe_card.line.color.rgb = BORDER_COLOR
    tf = fe_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FRONTEND SPA"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    fe_techs = [
        "React 18 + Vite: Khởi động tức thì, HMR siêu tốc.",
        "Tailwind CSS v3: Thiết kế UI đồng nhất, responsive.",
        "Lucide Icons: Bộ icon giao diện hiện đại.",
        "Axios Client: Quản lý interceptor token tự động.",
        "Dark / Light Mode: Trải nghiệm người dùng cao cấp."
    ]
    for it in fe_techs:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(10)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # Col 2: Backend
    be_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + col_w + gap), Inches(1.5), col_w, Inches(5.3))
    be_card.fill.solid()
    be_card.fill.fore_color.rgb = WHITE
    be_card.line.color.rgb = BORDER_COLOR
    tf = be_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "BACKEND CORE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    be_techs = [
        "Java 21 LTS + Spring Boot 3.x: Chuẩn công nghiệp.",
        "Spring Security 6: Stateless JWT Filter chain.",
        "Jakarta Bean Validation: Validate dữ liệu đầu vào.",
        "Global Exception Handler: API error chuẩn RFC 7807.",
        "Swagger / OpenAPI 3: Tài liệu API trực quan."
    ]
    for it in be_techs:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(10)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # Col 3: Database 3NF
    db_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + 2 * (col_w + gap)), Inches(1.5), col_w, Inches(5.3))
    db_card.fill.solid()
    db_card.fill.fore_color.rgb = WHITE
    db_card.line.color.rgb = BORDER_COLOR
    tf = db_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CSDL 13 BẢNG 3NF (POSTGRESQL)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GREEN_ACCENT
    db_techs = [
        "Flyway Migration (V1 - V8): Tự động hóa DDL/Seed.",
        "Phân hệ Người dùng: users, roles, permissions.",
        "Phân hệ Sảnh tiệc: venues, halls, hall_pricing (ca sáng/tối).",
        "Phân hệ Sự kiện: events, event_services, customers.",
        "Phân hệ Tài chính: menus, menu_items, services, payments."
    ]
    for it in db_techs:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(10)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 8: Live Product Demo (50% Milestone)
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Trình Diễn Sản Phẩm Trực Tiếp (Nghiệm Thu Mốc 50%)")

    # 4 Demo Step Cards
    demo_steps = [
        ("1. ĐĂNG NHẬP & BẢO MẬT", "Tài khoản: admin@evmanager.vn\nXác thực JWT Token, lưu trữ phiên an toàn, phân quyền Admin & Staff linh hoạt.", BLUE_ACCENT),
        ("2. DASHBOARD ĐIỀU HÀNH", "Hiển thị 4 thẻ KPI doanh thu & sảnh.\nBiểu đồ doanh thu trực quan, danh sách sự kiện gần nhất với trạng thái đã cọc/hoàn tất.", NAVY),
        ("3. QUẢN LÝ SẢNH & KHÁCH", "Tra cứu danh mục 4 sảnh tiệc lớn.\nLọc theo số lượng bàn, diện tích.\nQuản lý hồ sơ khách hàng và lịch sử tiệc cưới.", GREEN_ACCENT),
        ("4. BỘ TÍNH GIÁ TIỆC TỰ ĐỘNG", "Pricing Calculator tức thì trong 10s.\nTính giá sảnh theo ca sáng/tối, đơn giá thực đơn, phụ thu dịch vụ và tỷ lệ cọc 30%.", GOLD)
    ]

    for idx, (st_title, st_desc, st_color) in enumerate(demo_steps):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.95)
        y = Inches(1.5 + row * 2.7)
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.75), Inches(2.45))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = st_title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = st_color

        p = tf.add_paragraph()
        p.text = st_desc
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 9: Lessons Learned & Risk Management
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Bài Học Kinh Nghiệm & Kiểm Soát Rủi Ro Giai Đoạn 2")

    # Left: Lessons
    les_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.3))
    les_card.fill.solid()
    les_card.fill.fore_color.rgb = WHITE
    les_card.line.color.rgb = BORDER_COLOR
    tf = les_card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "BÀI HỌC KINH NGHIỆM ĐÃ RÚT RA"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT
    p.space_after = Pt(10)

    lessons = [
        "Chuẩn hóa dữ liệu sớm mang lại giá trị bền vững: Dành thêm thời gian làm CSDL 13 bảng 3NF giúp Backend và Frontend phát triển rất nhanh ở các sprint sau.",
        "Thiết lập API Contract trước khi code: Tách biệt Mock server giúp nhóm Frontend không bị block khi Backend đang hoàn thiện cấu hình security.",
        "Kỷ luật kiểm thử tự động: Chạy unit test trước khi merge PR giúp nhánh develop luôn hoạt động ổn định và sẵn sàng demo bất kỳ lúc nào."
    ]
    for les in lessons:
        p = tf.add_paragraph()
        p.text = "• " + les
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(10)

    # Right: Risks & Mitigation
    risk_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
    risk_card.fill.solid()
    risk_card.fill.fore_color.rgb = WHITE
    risk_card.line.color.rgb = BORDER_COLOR
    tf = risk_card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "KIỂM SOÁT RỦI RO CHO GIAI ĐOẠN 2 (TUẦN 4 - 8)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)

    risks = [
        "Ngăn ngừa phình phạm vi (Scope Creep): Mọi góp ý mới sau buổi demo giữa kỳ bắt buộc phải qua quy trình Change Request (CR-01 ở PM-21).",
        "Kiểm soát chi phí nhân sự: Tận dụng tối đa component UI và template CRUD có sẵn để nâng cao CPI lên >= 1.05.",
        "Rủi ro bảo mật & toàn vẹn dữ liệu: Áp dụng cơ chế @Transactional cho mọi thao tác thanh toán tiền cọc tiệc cưới.",
        "Rủi ro môi trường triển khai: Đóng gói Docker Compose chuẩn hóa để môi trường Staging chạy y hệt môi trường máy phát triển."
    ]
    for rsk in risks:
        p = tf.add_paragraph()
        p.text = "✔ " + rsk
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 10: Sprint 2 Roadmap & Q&A
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Kế Hoạch Nửa Sau Dự Án & Phiên Hỏi Đáp (Q&A)")

    # Roadmap box
    rd_box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(3.2))
    rd_box.fill.solid()
    rd_box.fill.fore_color.rgb = LIGHT_BG
    rd_box.line.color.rgb = BORDER_COLOR
    tf = rd_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "LỘ TRÌNH TRIỂN KHAI SPRINT 2 (TUẦN 4 ĐẾN TUẦN 8)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)

    milestones = [
        ("Tuần 4 - 5 (Sprint 2 Core):", "Hoàn thiện Module Hợp đồng tiệc cưới, Sơ đồ xếp bàn tiệc (Table Layout) và Thanh toán nhiều đợt."),
        ("Tuần 6 (Integration & Testing):", "Tích hợp toàn diện Backend & Frontend, Kiểm thử tự động E2E, Kiểm thử tải (Load Testing)."),
        ("Tuần 7 (Staging & Security Audit):", "Triển khai hệ thống lên môi trường Staging Docker, kiểm tra lỗ hổng bảo mật."),
        ("Tuần 8 (Final Release & Handover):", "Nghiệm thu toàn bộ 100% tính năng, bàn giao tài liệu kỹ thuật và đóng dự án.")
    ]
    for m_time, m_act in milestones:
        p = tf.add_paragraph()
        p.text = "▶ " + m_time + " " + m_act
        p.font.size = Pt(11)
        p.font.color.rgb = DARK
        p.space_before = Pt(6)

    # Q&A Banner
    qa_card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.9), Inches(11.7), Inches(2.0))
    qa_card.fill.solid()
    qa_card.fill.fore_color.rgb = NAVY
    qa_card.line.color.rgb = NAVY
    tf = qa_card.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "XIN CHÂN THÀNH CẢM ƠN HỘI ĐỒNG ĐÁNH GIÁ!"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Nhóm phát triển EVManager sẵn sàng tiếp nhận các câu hỏi và ý kiến đóng góp."
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(8)

    # Save
    output_path = os.path.join(os.path.dirname(__file__), "Midterm_Presentation_EVManager.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
