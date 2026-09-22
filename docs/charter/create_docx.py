import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Add a title
title = doc.add_heading('Tôn chỉ Dự án (Project Charter)', 1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Project title
p = doc.add_paragraph()
p.add_run('Tên dự án (Project Title): ').bold = True
p.add_run('EVManager — Hệ thống quản lý Trung tâm Hội nghị & Tiệc cưới (Mã: EVM-2026)')

# Dates
p = doc.add_paragraph()
p.add_run('Ngày bắt đầu (Project Start Date): ').bold = True
p.add_run('07/09/2026\t\t\t\t')
p.add_run('Ngày kết thúc (Projected Finish Date): ').bold = True
p.add_run('01/11/2026')

# Budget
p = doc.add_paragraph()
p.add_run('Thông tin về Kinh phí (Budget Information):\n').bold = True
p.add_run('Dự án được triển khai theo mô hình phát triển phần mềm theo yêu cầu trong 8 tuần:\n')
p.add_run('- Chi phí nhân sự (5 người x 20h/tuần x 8 tuần = 800 giờ x 100.000đ/giờ): 80.000.000 VNĐ\n')
p.add_run('- Chi phí hạ tầng (Cloud Server, Database, Domain, SSL): 3.000.000 VNĐ\n')
p.add_run('- Chi phí công cụ & Quản lý (Bản quyền phần mềm, Tools): 2.000.000 VNĐ\n')
p.add_run('- Quỹ dự phòng rủi ro (Contingency Reserve - 10%): 8.500.000 VNĐ\n')
p.add_run('=> Tổng ngân sách dự toán (Estimated Budget): 93.500.000 VNĐ.')

# PM
p = doc.add_paragraph()
p.add_run('GĐ Dự án (Project Manager): ').bold = True
p.add_run('Hoàng, ĐT: 09xx.xxx.xxx, Email: hoang@evmanager.local')

# Objectives
p = doc.add_paragraph()
p.add_run('Mục tiêu dự án (Project Objectives):\n').bold = True
p.add_run('Xây dựng hệ thống quản lý chuyên biệt cho Trung tâm Hội nghị & Tiệc cưới nhằm số hóa toàn bộ quy trình vận hành: Tiếp nhận khách -> Tư vấn chọn sảnh/thực đơn -> Báo giá -> Hợp đồng & Đặt cọc -> Điều phối sự kiện -> Thanh toán.\n')
p.add_run('Hệ thống hỗ trợ 4 phân hệ (Quản trị, Sales/Tư vấn, Điều phối/Bếp, Khách hàng), tự động cảnh báo xung đột lịch sảnh tiệc, cung cấp báo cáo doanh thu real-time. Đảm bảo hệ thống đạt hiệu năng cao, vận hành ổn định và bảo mật dữ liệu khách hàng.')

# Approach
p = doc.add_paragraph()
p.add_run('Cách tiếp cận (Approach):\n').bold = True
p.add_run('1. Lấy bài toán thực tế: Khảo sát quy trình vận hành của trung tâm tiệc cưới để phân tích nghiệp vụ lõi.\n')
p.add_run('2. Phương pháp Agile/Scrum: Phát triển trong 8 tuần, chia các Sprint 2 tuần để liên tục nghiệm thu tính năng và nhận feedback từ End-user.\n')
p.add_run('3. Quy trình kỹ thuật: Phân vai trò rõ ràng (PM, BA, Frontend, Backend, QA). Quản lý version qua Git, áp dụng tự động kiểm thử (CI/CD) và triển khai bản Beta (UAT) vào tuần 7.')

# Roles
p = doc.add_paragraph()
p.add_run('Vai trò và Trách nhiệm (Roles and Responsibilities):').bold = True

table = doc.add_table(rows=1, cols=4)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Vai trò'
hdr_cells[1].text = 'Họ Tên'
hdr_cells[2].text = 'Tổ chức/Vị trí\nOrganization/Position'
hdr_cells[3].text = 'Liên hệ\n(Contact information)'

for cell in hdr_cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

roles = [
    ('Project Manager', 'Hoàng', 'Khối Phát triển phần mềm', 'hoang@evmanager.local'),
    ('Business Analyst', 'Hiển', 'Khối Giải pháp Doanh nghiệp', 'hien@evmanager.local'),
    ('Frontend Developer', 'Nhân', 'Khối Phát triển phần mềm', 'nhan@evmanager.local'),
    ('Backend/DevOps', 'Phúc', 'Khối Phát triển phần mềm', 'phuc@evmanager.local'),
    ('QA / Tester', 'Hậu', 'Khối Đảm bảo Chất lượng', 'hau@evmanager.local')
]

for role, name, pos, contact in roles:
    row_cells = table.add_row().cells
    row_cells[0].text = role
    row_cells[1].text = name
    row_cells[2].text = pos
    row_cells[3].text = contact

# Sign-off
p = doc.add_paragraph('\n')
p.add_run('Ký tên (Sign-off): ').bold = True
p.add_run('(Chữ ký của các bên liên quan để xác nhận tài liệu.)\n\n')
p.add_run('Hoàng\t\tHiển\t\tNhân\t\tPhúc\t\tHậu\n')
p.add_run('..................\t..................\t..................\t..................\t..................')

# Comments
p = doc.add_paragraph('\n')
p.add_run('Chú thích (Comments): ').bold = True
p.add_run('Các tài liệu đặc tả yêu cầu (SRS), kế hoạch dự án (PMP) và thiết kế chi tiết được đính kèm ở thư mục /docs thuộc hệ thống quản trị dự án nội bộ.')

doc.save('d:\\EVManager\\docs\\charter\\Project_Charter_EVManager_Professional.docx')
