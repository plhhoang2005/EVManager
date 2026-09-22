import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Add a title
title = doc.add_heading('PHÁT BIỂU VỀ PHẠM VI - Scope Statement', 1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

table = doc.add_table(rows=6, cols=1)
table.style = 'Table Grid'

# Cell 0
cell_0 = table.cell(0,0)
p0 = cell_0.paragraphs[0]
p0.add_run('Tên dự án (Project Title): ').bold = True
p0.add_run('EVManager — Hệ thống quản lý Trung tâm Hội nghị & Tiệc cưới\n')
p0.add_run('Ngày (Date): ').bold = True
p0.add_run('12/09/2026\t\t\t')
p0.add_run('Người viết (Prepared by): ').bold = True
p0.add_run('Hoàng (PM)')

# Cell 1
cell_1 = table.cell(1,0)
p1 = cell_1.paragraphs[0]
p1.add_run('Lý Giải về dự án (Project Justification):\n').bold = True
p1.add_run('Hiện tại, việc quản lý đặt tiệc, sảnh hội nghị và các dịch vụ đi kèm bằng phương pháp thủ công dễ dẫn đến sai sót (trùng lịch, thất thoát thông tin). Dự án EVManager được triển khai nhằm số hóa toàn diện quy trình vận hành, tối ưu hóa việc quản lý sảnh trống, nâng cao tính chuyên nghiệp cho bộ phận Sales và giúp Ban Giám đốc kiểm soát chặt chẽ dòng tiền, hợp đồng.')

# Cell 2
cell_2 = table.cell(2,0)
p2 = cell_2.paragraphs[0]
p2.add_run('Các tính chất và yêu cầu của sản phẩm (Product Characteristics and Requirements):\n').bold = True
p2.add_run('1. Nền tảng: Hệ thống đa nền tảng (Web/Desktop) có giao diện trực quan, thân thiện (UI/UX), giúp tra cứu lịch sảnh nhanh chóng.\n')
p2.add_run('2. Bảo mật & Phân quyền: Cơ chế định danh và phân quyền (RBAC) nghiêm ngặt cho các phòng ban: Ban Giám đốc, Sales/Tư vấn, Điều phối/Kỹ thuật, Khách hàng.\n')
p2.add_run('3. Nghiệp vụ lõi: Đồng bộ luồng công việc: Chọn sảnh -> Chọn Menu/Dịch vụ -> Báo giá -> Hợp đồng/Đặt cọc -> Điều phối sự kiện -> Kế toán/Thanh toán.\n')
p2.add_run('4. Ràng buộc toàn vẹn: Tự động khóa sảnh và cảnh báo realtime để ngăn chặn việc đặt trùng sảnh sự kiện.')

# Cell 3
cell_3 = table.cell(3,0)
p3 = cell_3.paragraphs[0]
p3.add_run('Tổng kết về các sản phẩm chuyển giao của dự án (Summary of Project Deliverables)\n\n').bold = True
p3.add_run('Các kết quả liên quan đến quản lý dự án (Project management-related deliverables): ').bold = True
p3.add_run('Project Charter, Scope Statement, WBS, Gantt Chart, Biên bản họp (Meeting Minutes), Báo cáo tiến độ (Status reports), và bộ tài liệu nghiệm thu (Project Closure).')

# Cell 4
cell_4 = table.cell(4,0)
p4 = cell_4.paragraphs[0]
p4.add_run('Sản phẩm liên quan (Product-related deliverables): ').bold = True
p4.add_run('\n')
p4.add_run('1. Tài liệu đặc tả yêu cầu (SRS) và danh sách tính năng hệ thống.\n')
p4.add_run('2. Tài liệu thiết kế kiến trúc, CSDL (ERD) và thiết kế giao diện (UI/UX Prototype).\n')
p4.add_run('3. Mã nguồn phần mềm hoàn chỉnh và môi trường triển khai thực tế (Production).\n')
p4.add_run('4. Kịch bản kiểm thử (Test Cases), kịch bản UAT và Sổ tay Hướng dẫn sử dụng (User Manual).')

# Cell 5
cell_5 = table.cell(5,0)
p5 = cell_5.paragraphs[0]
p5.add_run('Các yêu cầu để đánh giá sự thành công của dự án (Project Success Criteria):\n').bold = True
p5.add_run('1. Về tiến độ: Triển khai và golive hệ thống đúng hạn trong thời gian 8 tuần theo cam kết Hợp đồng/Kế hoạch.\n')
p5.add_run('2. Về phạm vi: Hệ thống đáp ứng 100% các tính năng nghiệp vụ đã được khách hàng phê duyệt trong tài liệu đặc tả yêu cầu (SRS).\n')
p5.add_run('3. Về chất lượng: Vượt qua toàn bộ kịch bản nghiệm thu người dùng (UAT), hệ thống vận hành ổn định, không tồn tại lỗi nghiêm trọng (Severity 1, 2) trên môi trường Production.\n')
p5.add_run('4. Về quản trị: Minh bạch toàn bộ quá trình phát triển thông qua hệ thống quản lý source code (Github) và bàn giao đầy đủ mã nguồn, tài liệu hướng dẫn.')

doc.save('d:\\EVManager\\docs\\charter\\Scope_Statement_EVManager_Professional.docx')
